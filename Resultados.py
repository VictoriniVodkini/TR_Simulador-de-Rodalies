"""
Simulación de cierre de estaciones - Rodalies
-----------------------------------------------
Reutiliza el mismo grafo (pos, graf_base) y la misma lógica de coste
(Dijkstra con penalización de transbordo) que el simulador con interfaz,
pero sin Flet: para cada estación calcula qué pasaría con la red si esa
estación se cerrase, y vuelca los resultados a un Excel.
"""

import copy
import heapq
import time
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# 1. DATOS DEL GRAFO (idénticos al programa original, sin nada de interfaz)
# ---------------------------------------------------------------------------

pos = {
    "Sant Vicenç de Calders": (30, 650), "Calafell": (60, 650), "Segur de Calafell": (90, 650),
    "Cunit": (120, 650), "Cubelles": (150, 650), "Vilanova i la Geltrú": (180, 650),
    "Sitges": (210, 650), "Garraf": (240, 650), "Platja de Castelldefels": (270, 650),
    "Castelldefels": (300, 650), "Gavà": (330, 650), "Viladecans": (360, 650),
    "El Prat": (390, 650), "Bellvitge": (420, 650), "Aeroport": (390, 680),
    "Sant Andreu": (690, 560), "Montcada i Reixac": (730, 520), "La Llagosta": (750, 500),
    "Mollet-Sant Fost": (788, 500), "Montmeló": (826, 500), "Granollers-Centre": (864, 500),
    "Granollers-Nord": (902, 500), "Cardedeu": (940, 500), "Llinars del Vallés": (978, 500),
    "Palautordera": (1016, 500), "Sant Celoni": (1054, 500), "Gualba": (1092, 500),
    "Riells i Viabrea-Bred": (1130, 500), "Hostalric": (1168, 500), "Maçanet-Massanes": (1205, 500),
    "Sants": (480, 600), "Passeig de Gràcia": (530, 600), "Clot": (650, 600),
    "Hospitalet": (450, 600), "Cornellà": (400, 550), "Sant Joan Despí": (370, 510),
    "Sant Feliu": (370, 470), "Molins de Rei": (370, 430), "El Papiol": (370, 390),
    "Castellbisbal": (370, 350), "Martorell": (330, 350), "Rubí": (420, 350),
    "Sant Cugat": (520, 350), "Cerdanyola Universitat": (580, 350), "Santa Perpètua de Mogoda": (789, 350),
    "Plaça Catalunya": (530, 620), "Arc de Triomf": (550, 655), "La Sagrera-Meridiana": (650, 550),
    "Fabra i Puig": (680, 520), "Torre Baró-Vallbona": (710, 490), "Montcada Bifurcació": (740, 460),
    "Montcada-Manresa": (710, 430), "Montcada i Reixac-Santa Maria": (710, 400),
    "Cerdanyola del Vallès": (710, 370), "Barberà del Vallès": (710, 340), "Sabadell Sud": (710, 310),
    "Sabadell Centre": (710, 280), "Sabadell Nord": (710, 250),
    "Terrassa Est": (660, 250), "Terrassa Estació del Nord": (610, 250), "Sant Miquel de Gonteres": (560, 250),
    "Viladecavalls": (510, 250), "Vacarisses-Torreblanca": (460, 250), "Vacarisses": (410, 250),
    "Castellbell i el Vilar": (360, 250), "Sant Vicenç de Castellet": (310, 250), "Manresa": (260, 250),
    "Montcada-Ripollet": (770, 430), "La Florida": (820, 430), "Mollet-Santa Rosa": (870, 430),
    "Parets del Vallés": (920, 430), "Granollers-Canovelles": (970, 430), "Les Franqueses": (1020, 430),
    "La Garriga": (1040, 415), "Figaró": (1060, 400), "Sant Martí de Centelles": (1060, 375),
    "Centelles": (1060, 350), "Balenyà-Els Hostalets": (1060, 325), "Balenyà-Tona-Seva": (1060, 300),
    "Vic": (1060, 275), "Manlleu": (1060, 250), "Torrelló": (1060, 225), "Borgonyà": (1060, 200),
    "Sant Quirze de Besora": (1060, 175),
    "Sant Adrià de Besòs": (654, 650), "Badalona": (683, 650), "Montgat": (712, 650),
    "Montgat Nord": (741, 650), "El Masnou": (770, 650), "Ocata": (799, 650),
    "Premià de Mar": (828, 650), "Vilassar de Mar": (857, 650), "Cabrera de Mar": (886, 650),
    "Mataró": (915, 650), "Sant Andreu de Llavaneres": (944, 650), "Caldes d'Estrac": (973, 650),
    "Arenys de Mar": (1002, 650), "Canet de Mar": (1031, 650), "Sant Pol de Mar": (1060, 650),
    "Calella": (1089, 650), "Pineda de Mar": (1118, 650), "Santa Susanna": (1147, 650),
    "Malgrat de Mar": (1176, 650), "Blanes": (1205, 650), "Tordera": (1205, 580),
    "El Vendrell": (30, 590), "L'Arboç": (30, 530), "Els Monjos": (30, 470),
    "Vilafranca del Penedès": (30, 410), "La Granada": (80, 410), "Lavern-Subirats": (130, 410),
    "Sant Sadurní d'Anoia": (180, 410), "Gelida": (230, 410),
}

graf_base = {
    "Aeroport": [("El Prat", "R2N")],
    "Sant Vicenç de Calders": [("Calafell", "R2S"), ("El Vendrell", "R4")],
    "El Vendrell": [("Sant Vicenç de Calders", "R4"), ("L'Arboç", "R4")],
    "L'Arboç": [("El Vendrell", "R4"), ("Els Monjos", "R4")],
    "Els Monjos": [("L'Arboç", "R4"), ("Vilafranca del Penedès", "R4")],
    "Vilafranca del Penedès": [("Els Monjos", "R4"), ("La Granada", "R4")],
    "La Granada": [("Vilafranca del Penedès", "R4"), ("Lavern-Subirats", "R4")],
    "Lavern-Subirats": [("La Granada", "R4"), ("Sant Sadurní d'Anoia", "R4")],
    "Sant Sadurní d'Anoia": [("Lavern-Subirats", "R4"), ("Gelida", "R4")],
    "Gelida": [("Sant Sadurní d'Anoia", "R4"), ("Martorell", "R4")],
    "Calafell": [("Sant Vicenç de Calders", "R2S"), ("Segur de Calafell", "R2S")],
    "Segur de Calafell": [("Calafell", "R2S"), ("Cunit", "R2S")],
    "Cunit": [("Segur de Calafell", "R2S"), ("Cubelles", "R2S")],
    "Cubelles": [("Cunit", "R2S"), ("Vilanova i la Geltrú", "R2S")],
    "Vilanova i la Geltrú": [("Cubelles", "R2S"), ("Sitges", "R2S")],
    "Sitges": [("Vilanova i la Geltrú", "R2S"), ("Garraf", "R2S")],
    "Garraf": [("Sitges", "R2S"), ("Platja de Castelldefels", "R2S")],
    "Platja de Castelldefels": [("Garraf", "R2S"), ("Castelldefels", "R2S")],
    "Castelldefels": [("Gavà", "R2"), ("Platja de Castelldefels", "R2S"), ("Gavà", "R2S")],
    "Gavà": [("Castelldefels", "R2"), ("Viladecans", "R2"), ("Castelldefels", "R2S"), ("Viladecans", "R2S")],
    "Viladecans": [("Gavà", "R2"), ("El Prat", "R2"), ("Gavà", "R2S"), ("El Prat", "R2S")],
    "El Prat": [("Viladecans", "R2"), ("Bellvitge", "R2"), ("Aeroport", "R2N"), ("Bellvitge", "R2N"),
                ("Viladecans", "R2S"), ("Bellvitge", "R2S")],
    "Bellvitge": [("El Prat", "R2"), ("Sants", "R2"), ("El Prat", "R2N"), ("Sants", "R2N"),
                  ("El Prat", "R2S"), ("Sants", "R2S")],
    "Sants": [("Bellvitge", "R2"), ("Passeig de Gràcia", "R2"), ("Bellvitge", "R2N"), ("Passeig de Gràcia", "R2N"),
              ("Bellvitge", "R2S"), ("Passeig de Gràcia", "R2S"), ("Hospitalet", "R1"), ("Plaça Catalunya", "R1"),
              ("Hospitalet", "R3"), ("Plaça Catalunya", "R3"), ("Hospitalet", "R4"), ("Plaça Catalunya", "R4")],
    "Passeig de Gràcia": [("Sants", "R2"), ("Clot", "R2"), ("Sants", "R2N"), ("Clot", "R2N"),
                           ("Sants", "R2S"), ("Clot", "R2S")],
    "Clot": [("Passeig de Gràcia", "R2"), ("Sant Andreu", "R2"), ("Arc de Triomf", "R1"),
             ("Passeig de Gràcia", "R2N"), ("Sant Andreu", "R2N")],
    "Sant Andreu": [("Montcada i Reixac", "R2"), ("Clot", "R2"), ("Montcada i Reixac", "R2N"), ("Clot", "R2N")],
    "Montcada i Reixac": [("La Llagosta", "R2"), ("Sant Andreu", "R2"), ("La Llagosta", "R2N"), ("Sant Andreu", "R2N")],
    "La Llagosta": [("Montcada i Reixac", "R2"), ("Mollet-Sant Fost", "R2"), ("Montcada i Reixac", "R2N"), ("Mollet-Sant Fost", "R2N")],
    "Mollet-Sant Fost": [("La Llagosta", "R2"), ("Montmeló", "R2"), ("La Llagosta", "R2N"), ("Montmeló", "R2N"),
                          ("Santa Perpètua de Mogoda", "R8"), ("Montmeló", "R8")],
    "Montmeló": [("Granollers-Centre", "R2"), ("Mollet-Sant Fost", "R2"), ("Granollers-Centre", "R2N"),
                 ("Mollet-Sant Fost", "R2N"), ("Granollers-Centre", "R8"), ("Mollet-Sant Fost", "R8")],
    "Granollers-Centre": [("Montmeló", "R2"), ("Montmeló", "R2N"), ("Granollers-Nord", "R2N"), ("Montmeló", "R8")],
    "Granollers-Nord": [("Granollers-Centre", "R2N"), ("Cardedeu", "R2N")],
    "Cardedeu": [("Granollers-Nord", "R2N"), ("Llinars del Vallés", "R2N")],
    "Llinars del Vallés": [("Palautordera", "R2N"), ("Cardedeu", "R2N")],
    "Palautordera": [("Sant Celoni", "R2N"), ("Llinars del Vallés", "R2N")],
    "Sant Celoni": [("Palautordera", "R2N"), ("Gualba", "R2N")],
    "Gualba": [("Sant Celoni", "R2N"), ("Riells i Viabrea-Bred", "R2N")],
    "Riells i Viabrea-Bred": [("Hostalric", "R2N"), ("Gualba", "R2N")],
    "Hostalric": [("Maçanet-Massanes", "R2N"), ("Riells i Viabrea-Bred", "R2N")],
    "Maçanet-Massanes": [("Tordera", "R1"), ("Hostalric", "R2N")],
    "Santa Perpètua de Mogoda": [("Mollet-Sant Fost", "R8"), ("Cerdanyola Universitat", "R8")],
    "Cerdanyola Universitat": [("Sant Cugat", "R8"), ("Santa Perpètua de Mogoda", "R8"), ("Cerdanyola del Vallès", "R7")],
    "Sant Cugat": [("Rubí", "R8"), ("Cerdanyola Universitat", "R8")],
    "Rubí": [("Castellbisbal", "R8"), ("Sant Cugat", "R8")],
    "Martorell": [("Gelida", "R4"), ("Castellbisbal", "R4"), ("Castellbisbal", "R8")],
    "Castellbisbal": [("Martorell", "R4"), ("El Papiol", "R4"), ("Martorell", "R8"), ("Rubí", "R8")],
    "El Papiol": [("Molins de Rei", "R4"), ("Castellbisbal", "R4")],
    "Molins de Rei": [("Sant Feliu", "R1"), ("Sant Feliu", "R4"), ("El Papiol", "R4")],
    "Sant Feliu": [("Sant Joan Despí", "R1"), ("Molins de Rei", "R1"), ("Sant Joan Despí", "R4"), ("Molins de Rei", "R4")],
    "Sant Joan Despí": [("Cornellà", "R1"), ("Sant Feliu", "R1"), ("Cornellà", "R4"), ("Sant Feliu", "R4")],
    "Cornellà": [("Hospitalet", "R1"), ("Sant Joan Despí", "R1"), ("Hospitalet", "R4"), ("Sant Joan Despí", "R4")],
    "Hospitalet": [("Sants", "R1"), ("Cornellà", "R1"), ("Sants", "R3"), ("Sants", "R4"), ("Cornellà", "R4")],
    "Plaça Catalunya": [("Sants", "R1"), ("Arc de Triomf", "R1"), ("Sants", "R3"), ("Arc de Triomf", "R3"),
                         ("Sants", "R4"), ("Arc de Triomf", "R4")],
    "Arc de Triomf": [("Plaça Catalunya", "R1"), ("Clot", "R1"), ("Plaça Catalunya", "R3"), ("La Sagrera-Meridiana", "R3"),
                       ("Plaça Catalunya", "R4"), ("La Sagrera-Meridiana", "R4")],
    "La Sagrera-Meridiana": [("Arc de Triomf", "R3"), ("Fabra i Puig", "R3"), ("Arc de Triomf", "R4"), ("Fabra i Puig", "R4")],
    "Fabra i Puig": [("La Sagrera-Meridiana", "R3"), ("Torre Baró-Vallbona", "R3"), ("La Sagrera-Meridiana", "R4"),
                      ("Torre Baró-Vallbona", "R4"), ("Torre Baró-Vallbona", "R7")],
    "Torre Baró-Vallbona": [("Fabra i Puig", "R3"), ("Montcada Bifurcació", "R3"), ("Fabra i Puig", "R4"),
                             ("Montcada Bifurcació", "R4"), ("Fabra i Puig", "R7"), ("Montcada Bifurcació", "R7")],
    "Montcada Bifurcació": [("Torre Baró-Vallbona", "R3"), ("Montcada-Ripollet", "R3"), ("Torre Baró-Vallbona", "R4"),
                             ("Montcada-Manresa", "R4"), ("Torre Baró-Vallbona", "R7"), ("Montcada-Manresa", "R7")],
    "Montcada-Manresa": [("Montcada Bifurcació", "R4"), ("Montcada i Reixac-Santa Maria", "R4"),
                          ("Montcada Bifurcació", "R7"), ("Montcada i Reixac-Santa Maria", "R7")],
    "Montcada i Reixac-Santa Maria": [("Montcada-Manresa", "R4"), ("Cerdanyola del Vallès", "R4"),
                                       ("Montcada-Manresa", "R7"), ("Cerdanyola del Vallès", "R7")],
    "Cerdanyola del Vallès": [("Montcada i Reixac-Santa Maria", "R4"), ("Barberà del Vallès", "R4"),
                               ("Montcada i Reixac-Santa Maria", "R7"), ("Cerdanyola Universitat", "R7")],
    "Barberà del Vallès": [("Cerdanyola del Vallès", "R4"), ("Sabadell Sud", "R4")],
    "Sabadell Sud": [("Barberà del Vallès", "R4"), ("Sabadell Centre", "R4")],
    "Sabadell Centre": [("Sabadell Sud", "R4"), ("Sabadell Nord", "R4")],
    "Sabadell Nord": [("Sabadell Centre", "R4"), ("Terrassa Est", "R4")],
    "Terrassa Est": [("Sabadell Nord", "R4"), ("Terrassa Estació del Nord", "R4")],
    "Terrassa Estació del Nord": [("Terrassa Est", "R4"), ("Sant Miquel de Gonteres", "R4")],
    "Sant Miquel de Gonteres": [("Terrassa Estació del Nord", "R4"), ("Viladecavalls", "R4")],
    "Viladecavalls": [("Sant Miquel de Gonteres", "R4"), ("Vacarisses-Torreblanca", "R4")],
    "Vacarisses-Torreblanca": [("Viladecavalls", "R4"), ("Vacarisses", "R4")],
    "Vacarisses": [("Vacarisses-Torreblanca", "R4"), ("Castellbell i el Vilar", "R4")],
    "Castellbell i el Vilar": [("Vacarisses", "R4"), ("Sant Vicenç de Castellet", "R4")],
    "Sant Vicenç de Castellet": [("Castellbell i el Vilar", "R4"), ("Manresa", "R4")],
    "Manresa": [("Sant Vicenç de Castellet", "R4")],
    "Montcada-Ripollet": [("Montcada Bifurcació", "R3"), ("La Florida", "R3")],
    "La Florida": [("Montcada-Ripollet", "R3"), ("Mollet-Santa Rosa", "R3")],
    "Mollet-Santa Rosa": [("Parets del Vallés", "R3"), ("La Florida", "R3")],
    "Parets del Vallés": [("Mollet-Santa Rosa", "R3"), ("Granollers-Canovelles", "R3")],
    "Granollers-Canovelles": [("Parets del Vallés", "R3"), ("Les Franqueses", "R3")],
    "Les Franqueses": [("La Garriga", "R3"), ("Granollers-Canovelles", "R3")],
    "La Garriga": [("Figaró", "R3"), ("Les Franqueses", "R3")],
    "Figaró": [("La Garriga", "R3"), ("Sant Martí de Centelles", "R3")],
    "Sant Martí de Centelles": [("Figaró", "R3"), ("Centelles", "R3")],
    "Centelles": [("Balenyà-Els Hostalets", "R3"), ("Sant Martí de Centelles", "R3")],
    "Balenyà-Els Hostalets": [("Balenyà-Tona-Seva", "R3"), ("Centelles", "R3")],
    "Balenyà-Tona-Seva": [("Balenyà-Els Hostalets", "R3"), ("Vic", "R3")],
    "Vic": [("Balenyà-Tona-Seva", "R3"), ("Manlleu", "R3")],
    "Manlleu": [("Torrelló", "R3"), ("Vic", "R3")],
    "Torrelló": [("Borgonyà", "R3"), ("Manlleu", "R3")],
    "Borgonyà": [("Torrelló", "R3"), ("Sant Quirze de Besora", "R3")],
    "Sant Quirze de Besora": [("Borgonyà", "R3")],
    "Sant Adrià de Besòs": [("Clot", "R1"), ("Badalona", "R1")],
    "Badalona": [("Sant Adrià de Besòs", "R1"), ("Montgat", "R1")],
    "Montgat": [("Montgat Nord", "R1"), ("Badalona", "R1")],
    "Montgat Nord": [("El Masnou", "R1"), ("Montgat", "R1")],
    "El Masnou": [("Montgat Nord", "R1"), ("Ocata", "R1")],
    "Ocata": [("El Masnou", "R1"), ("Premià de Mar", "R1")],
    "Premià de Mar": [("Vilassar de Mar", "R1"), ("Ocata", "R1")],
    "Vilassar de Mar": [("Cabrera de Mar", "R1"), ("Premià de Mar", "R1")],
    "Cabrera de Mar": [("Vilassar de Mar", "R1"), ("Mataró", "R1")],
    "Mataró": [("Cabrera de Mar", "R1"), ("Sant Andreu de Llavaneres", "R1")],
    "Sant Andreu de Llavaneres": [("Caldes d'Estrac", "R1"), ("Mataró", "R1")],
    "Caldes d'Estrac": [("Arenys de Mar", "R1"), ("Sant Andreu de Llavaneres", "R1")],
    "Arenys de Mar": [("Caldes d'Estrac", "R1"), ("Canet de Mar", "R1")],
    "Canet de Mar": [("Arenys de Mar", "R1"), ("Sant Pol de Mar", "R1")],
    "Sant Pol de Mar": [("Calella", "R1"), ("Canet de Mar", "R1")],
    "Calella": [("Pineda de Mar", "R1"), ("Sant Pol de Mar", "R1")],
    "Pineda de Mar": [("Calella", "R1"), ("Santa Susanna", "R1")],
    "Santa Susanna": [("Pineda de Mar", "R1"), ("Malgrat de Mar", "R1")],
    "Malgrat de Mar": [("Blanes", "R1"), ("Santa Susanna", "R1")],
    "Blanes": [("Tordera", "R1"), ("Malgrat de Mar", "R1")],
    "Tordera": [("Blanes", "R1"), ("Maçanet-Massanes", "R1")],
}

# ---------------------------------------------------------------------------
# 2. LÓGICA DE COSTE (idéntica a calcular_coste_simulacion, sin depender de "state")
# ---------------------------------------------------------------------------

def costes_desde_origen(g, estaciones):
    """Coste mínimo desde cada estación de origen hasta todas las demás,
    con Dijkstra + penalización de 5 por transbordo de línea."""
    resultado = {}
    for inicio in estaciones:
        cola = [(0, inicio, None)]
        visitados = {}
        mejores = {}
        while cola:
            coste, nodo, linea_actual = heapq.heappop(cola)
            estado = (nodo, linea_actual)
            if estado in visitados and visitados[estado] <= coste:
                continue
            visitados[estado] = coste
            if nodo not in mejores or coste < mejores[nodo]:
                mejores[nodo] = coste
            for vecino, linea in g.get(nodo, []):
                nuevo_coste = coste + 1
                if linea_actual is not None and linea != linea_actual:
                    nuevo_coste += 5
                heapq.heappush(cola, (nuevo_coste, vecino, linea))
        resultado[inicio] = mejores
    return resultado


def cerrar_estacion(grafo, estacion):
    """Devuelve una copia del grafo con 'estacion' cerrada
    (igual que hace on_click_der en la versión con interfaz)."""
    nuevo = copy.deepcopy(grafo)
    for v in nuevo:
        nuevo[v] = [(d, l) for d, l in nuevo[v] if d != estacion]
    nuevo[estacion] = []
    return nuevo


def calcular_coste_simulacion(grafo, cerradas, costes_base):
    """Misma métrica que en el programa original, pero recibiendo
    'cerradas' como parámetro en lugar de leerlo de state[]."""
    costes_actuales = costes_desde_origen(grafo, pos.keys())

    estaciones = [e for e in pos if e not in cerradas]

    suma_costes_actuales = 0
    suma_costes_base_posibles = 0
    rutas_posibles = 0
    rutas_imposibles = 0
    caminos_afectados = 0

    for i in range(len(estaciones)):
        for j in range(i + 1, len(estaciones)):
            origen = estaciones[i]
            destino = estaciones[j]

            coste_base = costes_base.get(origen, {}).get(destino)
            coste_actual = costes_actuales.get(origen, {}).get(destino)

            if coste_actual is None:
                rutas_imposibles += 1
            else:
                rutas_posibles += 1
                suma_costes_actuales += coste_actual
                if coste_base is not None:
                    suma_costes_base_posibles += coste_base
                    if coste_actual > coste_base:
                        caminos_afectados += 1

    coste_medio = suma_costes_actuales / rutas_posibles if rutas_posibles else 0
    coste_base_medio = suma_costes_base_posibles / rutas_posibles if rutas_posibles else 0
    incremento = coste_medio - coste_base_medio
    incremento_porcentaje = (incremento / coste_base_medio * 100) if coste_base_medio else 0

    total_caminos = len(estaciones) * (len(estaciones) - 1) // 2
    porcentaje_afectados = (caminos_afectados / total_caminos * 100) if total_caminos else 0

    return {
        "rutas_posibles": rutas_posibles,
        "rutas_imposibles": rutas_imposibles,
        "caminos_afectados": caminos_afectados,
        "porcentaje_afectados": porcentaje_afectados,
        "coste_medio": coste_medio,
        "incremento": incremento,
        "incremento_porcentaje": incremento_porcentaje,
    }


# ---------------------------------------------------------------------------
# 3. SIMULACIÓN: cerrar cada estación, una a una, y medir el impacto
# ---------------------------------------------------------------------------

def simular_todas_las_estaciones():
    print("Calculando costes base de la xarxa completa...")
    costes_base = costes_desde_origen(graf_base, pos.keys())

    resultados = []
    estaciones = sorted(pos.keys())
    total = len(estaciones)

    for idx, estacion in enumerate(estaciones, start=1):
        t0 = time.time()
        grafo_cerrado = cerrar_estacion(graf_base, estacion)
        metricas = calcular_coste_simulacion(grafo_cerrado, {estacion}, costes_base)
        metricas["estacion"] = estacion
        resultados.append(metricas)
        print(f"  [{idx}/{total}] {estacion} -> "
              f"impossibles={metricas['rutas_imposibles']}, "
              f"afectats={metricas['porcentaje_afectados']:.1f}%, "
              f"increment={metricas['incremento_porcentaje']:+.1f}% "
              f"({time.time()-t0:.2f}s)")

    return resultados


# ---------------------------------------------------------------------------
# 4. EXPORTAR A EXCEL
# ---------------------------------------------------------------------------

def exportar_excel(resultados, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Impacte tancaments"

    headers = [
        "Estació tancada",
        "Rutes possibles",
        "Rutes impossibles",
        "Camins afectats",
        "% Camins afectats",
        "Cost mitjà (amb tancament)",
        "Increment mitjà",
        "Increment %",
    ]
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="045C1B")
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", wrap_text=True)

    # Ordenamos las estaciones por impacto (más rutas impossibles primero,
    # luego mayor incremento porcentual) para que las más críticas queden arriba.
    resultados_ordenados = sorted(
        resultados,
        key=lambda r: (r["rutas_imposibles"], r["incremento_porcentaje"]),
        reverse=True,
    )

    for row, r in enumerate(resultados_ordenados, start=2):
        ws.cell(row=row, column=1, value=r["estacion"]).font = Font(name="Arial")
        ws.cell(row=row, column=2, value=r["rutas_posibles"]).font = Font(name="Arial")
        ws.cell(row=row, column=3, value=r["rutas_imposibles"]).font = Font(name="Arial")
        ws.cell(row=row, column=4, value=r["caminos_afectados"]).font = Font(name="Arial")
        c5 = ws.cell(row=row, column=5, value=round(r["porcentaje_afectados"] / 100, 4))
        c5.number_format = "0.00%"
        c5.font = Font(name="Arial")
        c6 = ws.cell(row=row, column=6, value=round(r["coste_medio"], 3))
        c6.font = Font(name="Arial")
        c7 = ws.cell(row=row, column=7, value=round(r["incremento"], 3))
        c7.font = Font(name="Arial")
        c8 = ws.cell(row=row, column=8, value=round(r["incremento_porcentaje"] / 100, 4))
        c8.number_format = "0.00%"
        c8.font = Font(name="Arial")

        # Resaltar en rojo suave las estaciones que dejan rutas imposibles
        if r["rutas_imposibles"] > 0:
            fill_alerta = PatternFill("solid", fgColor="FBE2E2")
            for col in range(1, 9):
                ws.cell(row=row, column=col).fill = fill_alerta

    # Anchuras de columna
    anchos = [32, 15, 17, 15, 16, 22, 16, 13]
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    ws.freeze_panes = "A2"

    # Nota explicativa al final
    nota_row = len(resultados_ordenados) + 3
    nota = ws.cell(
        row=nota_row, column=1,
        value=("Nota: 'Rutes impossibles' = parells d'estacions que queden totalment desconnectats "
               "en tancar aquesta estació. 'Cost mitjà' inclou penalització de +5 per transbordament, "
               "igual que al simulador. Càlcul generat des del mateix graf i lògica de Dijkstra del programa original.")
    )
    nota.font = Font(name="Arial", italic=True, size=9, color="666666")
    ws.merge_cells(start_row=nota_row, start_column=1, end_row=nota_row, end_column=8)
    nota.alignment = Alignment(wrap_text=True)

    wb.save(path)


if __name__ == "__main__":
    resultados = simular_todas_las_estaciones()
    salida = r"C:\Users\victo\OneDrive\Escritorio\TR-Grafos\impacto_cierre_estaciones.xlsx"
    exportar_excel(resultados, salida)
    print(f"\nExcel generat a: {salida}")